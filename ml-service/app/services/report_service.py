from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle, Paragraph

from app.services.health_service import calculate_health_score
from app.services.insights_service import generate_offline_insights
from app.services.stats_service import get_numeric_stats


def build_report(records: list[dict], report_format: str, dataset_name: str = "dataset") -> tuple[bytes, str, str]:
    frame = pd.DataFrame(records)
    health = calculate_health_score(frame)
    stats = get_numeric_stats(frame)
    insights = generate_offline_insights(frame)
    report_format = report_format.upper()

    if report_format == "PDF":
        return _build_pdf(dataset_name, health, stats, insights), "application/pdf", f"{dataset_name}-report.pdf"
    if report_format == "XLSX":
        return _build_xlsx(frame, dataset_name, health, stats, insights), \
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"{dataset_name}-report.xlsx"
    raise ValueError("format must be PDF or XLSX")


def _build_pdf(dataset_name: str, health: dict, stats: dict, insights: dict) -> bytes:
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=letter, rightMargin=36, leftMargin=36,
                                 topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    content = [Paragraph(f"AI Data Analyst Report: {dataset_name}", styles["Title"]), Spacer(1, 12)]
    content.append(Paragraph(f"Health score: {health['score']}% ({health['grade']})", styles["Heading2"]))
    content.append(_table(["Metric", "Value"], [[key.replace("_", " ").title(), value]
                           for key, value in health.items() if key != "issues"]))
    content.append(Spacer(1, 12))
    content.append(Paragraph("Numeric statistics", styles["Heading2"]))
    stat_rows = [["Column", "Count", "Mean", "Median", "Min", "Max"]]
    for column, values in stats.items():
        stat_rows.append([column, values.get("count"), values.get("mean"), values.get("median"),
                          values.get("min"), values.get("max")])
    content.append(_table(stat_rows[0], stat_rows[1:]))
    content.append(Spacer(1, 12))
    content.append(Paragraph("Insights", styles["Heading2"]))
    for category in ("anomalies", "trends", "correlations"):
        content.append(Paragraph(category.title(), styles["Heading3"]))
        values = insights.get(category) or ["None detected."]
        content.extend(Paragraph(f"- {value}", styles["BodyText"]) for value in values)
    document.build(content)
    return output.getvalue()


def _table(headers: list, rows: list[list]) -> Table:
    table = Table([headers] + rows, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#243447")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]))
    return table


def _build_xlsx(frame: pd.DataFrame, dataset_name: str, health: dict, stats: dict, insights: dict) -> bytes:
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Dataset"
    for row in [list(frame.columns)] + frame.astype(object).where(pd.notna(frame), "").values.tolist():
        data_sheet.append(row)
    _style_header(data_sheet)

    health_sheet = workbook.create_sheet("Summary")
    health_sheet.append(["Health metric", "Value"])
    for key, value in health.items():
        if key != "issues":
            health_sheet.append([key, value])
    health_sheet.append([])
    health_sheet.append(["Insight category", "Finding"])
    for category in ("anomalies", "trends", "correlations"):
        for finding in insights.get(category) or ["None detected."]:
            health_sheet.append([category, finding])
    _style_header(health_sheet)

    stats_sheet = workbook.create_sheet("Statistics")
    stat_keys = ["column", "count", "mean", "median", "min", "max", "std", "skewness", "kurtosis"]
    stats_sheet.append(stat_keys)
    for column, values in stats.items():
        stats_sheet.append([column] + [values.get(key) for key in stat_keys[1:]])
    _style_header(stats_sheet)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _style_header(sheet) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = __import__("openpyxl").styles.PatternFill("solid", fgColor="243447")
    sheet.freeze_panes = "A2"
